#!/usr/bin/env python3
"""
Test suite for the Greyhound PDF ingestion system.

This module contains basic tests for the core functionality.
Run with: python -m pytest tests/
"""

import unittest
import sys
from pathlib import Path
import tempfile
import os

# Add src to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))


class TestPDFParsing(unittest.TestCase):
    """Tests for PDF parsing functionality."""
    
    def test_import_parse_pdf(self):
        """Test that parse_pdf module can be imported."""
        try:
            import parse_pdf
            self.assertIsNotNone(parse_pdf)
        except ImportError as e:
            self.fail(f"Failed to import parse_pdf: {e}")
    
    def test_filename_regex(self):
        """Test filename pattern matching."""
        import parse_pdf
        import re
        
        # Test valid filenames
        valid_names = [
            "SALE_2025-09-01.pdf",
            "RICH_2024-12-31.pdf",
            "CAPA_2025-01-15.pdf"
        ]
        
        for name in valid_names:
            match = parse_pdf.FNAME_RE.match(name)
            self.assertIsNotNone(match, f"Should match valid filename: {name}")
            self.assertEqual(len(match.groups()), 2)
        
        # Test invalid filenames
        invalid_names = [
            "SALE_2025-09-01.txt",
            "sale_2025-09-01.pdf",
            "SALE2025-09-01.pdf",
            "SALE_20250901.pdf"
        ]
        
        for name in invalid_names:
            match = parse_pdf.FNAME_RE.match(name)
            self.assertIsNone(match, f"Should not match invalid filename: {name}")


class TestFetchForms(unittest.TestCase):
    """Tests for PDF fetching functionality."""
    
    def test_import_fetch_forms(self):
        """Test that fetch_forms module can be imported."""
        try:
            import fetch_forms
            self.assertIsNotNone(fetch_forms)
        except ImportError as e:
            self.fail(f"Failed to import fetch_forms: {e}")
    
    def test_reject_patterns(self):
        """Test that sponsor/promotional PDFs are rejected."""
        import fetch_forms
        
        # These should be rejected
        reject_cases = [
            ("ladbrokes_special.pdf", "context"),
            ("club_form.pdf", "context"),
            ("million_dollar_chase.pdf", "context"),
        ]
        
        for url, context in reject_cases:
            result = fetch_forms._looks_like_form(url, context)
            self.assertFalse(result, f"Should reject: {url}")


class TestMastercontrol(unittest.TestCase):
    """Tests for main automation script."""
    
    def test_import_mastercontrol(self):
        """Test that mastercontrol module can be imported."""
        try:
            import mastercontrol
            self.assertIsNotNone(mastercontrol)
        except ImportError as e:
            self.fail(f"Failed to import mastercontrol: {e}")
    
    def test_build_probabilities_empty(self):
        """Test probability calculation with empty DataFrame."""
        import pandas as pd
        import mastercontrol
        
        empty_df = pd.DataFrame()
        result = mastercontrol.build_probabilities(empty_df)
        
        self.assertTrue(result.empty)
        self.assertListEqual(
            list(result.columns),
            ["track", "date", "race", "box", "runner", "prob_win"]
        )
    
    def test_build_probabilities_basic(self):
        """Test probability calculation with sample data."""
        import pandas as pd
        import mastercontrol
        
        # Create sample data: 1 race with 4 runners
        sample_data = pd.DataFrame([
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 1, "runner": "DOG A"},
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 2, "runner": "DOG B"},
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 3, "runner": "DOG C"},
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 4, "runner": "DOG D"},
        ])
        
        result = mastercontrol.build_probabilities(sample_data)
        
        self.assertEqual(len(result), 4)
        self.assertTrue(all(result["prob_win"] == 0.25))
        self.assertIn("prob_win", result.columns)


class TestDataIntegrity(unittest.TestCase):
    """Tests for data validation and integrity."""
    
    def test_directory_structure(self):
        """Test that required directories exist."""
        root = Path(__file__).resolve().parent.parent
        
        required_dirs = [
            root / "src",
            root / "data" / "input",
            root / "data" / "output",
            root / "tests",
        ]
        
        for dir_path in required_dirs:
            self.assertTrue(
                dir_path.exists(),
                f"Required directory should exist: {dir_path}"
            )
    
    def test_required_files(self):
        """Test that required files exist."""
        root = Path(__file__).resolve().parent.parent
        
        required_files = [
            root / "requirements.txt",
            root / "README.md",
            root / "LICENSE",
            root / ".gitignore",
        ]
        
        for file_path in required_files:
            self.assertTrue(
                file_path.exists(),
                f"Required file should exist: {file_path}"
            )


class TestExcelExport(unittest.TestCase):
    """Tests for Excel export functionality."""
    
    def test_import_export_to_excel(self):
        """Test that export_to_excel module can be imported."""
        try:
            import export_to_excel
            self.assertIsNotNone(export_to_excel)
        except ImportError as e:
            self.fail(f"Failed to import export_to_excel: {e}")
    
    def test_create_dog_summary_df(self):
        """Test Dog Summary DataFrame creation."""
        import pandas as pd
        import export_to_excel
        
        # Test with sample data
        sample_data = pd.DataFrame([
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 1, "runner": "DOG A"},
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 2, "runner": "DOG B"},
        ])
        
        result = export_to_excel.create_dog_summary_df(sample_data)
        
        self.assertEqual(len(result), 2)
        self.assertIn("Dog_Name", result.columns)
        self.assertIn("Race_No", result.columns)
        self.assertIn("Tab_No", result.columns)
        # Check for 31 columns as specified
        self.assertEqual(len(result.columns), 31)
    
    def test_create_race_history_df(self):
        """Test Race History DataFrame creation."""
        import pandas as pd
        import export_to_excel
        
        # Test with sample data
        sample_data = pd.DataFrame([
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 1, "runner": "DOG A"},
        ])
        
        result = export_to_excel.create_race_history_df(sample_data)
        
        self.assertEqual(len(result), 1)
        self.assertIn("Dog_Name", result.columns)
        self.assertIn("Hist_Date", result.columns)
        self.assertIn("Hist_Track", result.columns)
        # Check for 23 columns as specified
        self.assertEqual(len(result.columns), 23)
    
    def test_export_to_excel_creates_file(self):
        """Test that export_to_excel creates a valid Excel file."""
        import pandas as pd
        import export_to_excel
        from openpyxl import load_workbook
        
        # Create test data
        sample_data = pd.DataFrame([
            {"track": "TEST", "date": "2025-01-01", "race": 1, "box": 1, "runner": "DOG A"},
        ])
        
        summary_df = export_to_excel.create_dog_summary_df(sample_data)
        history_df = export_to_excel.create_race_history_df(sample_data)
        
        # Create temporary file
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
            temp_path = tmp.name
        
        try:
            # Export to Excel
            export_to_excel.export_to_excel(summary_df, history_df, temp_path)
            
            # Verify file was created
            self.assertTrue(os.path.exists(temp_path))
            
            # Verify Excel structure
            wb = load_workbook(temp_path)
            self.assertIn("Dog Summary", wb.sheetnames)
            self.assertIn("Race History Detail", wb.sheetnames)
            
            # Verify Dog Summary has correct columns
            ws_summary = wb["Dog Summary"]
            self.assertEqual(ws_summary.max_column, 31)
            
            # Verify Race History has correct columns
            ws_history = wb["Race History Detail"]
            self.assertEqual(ws_history.max_column, 23)
            
        finally:
            # Clean up
            if os.path.exists(temp_path):
                os.unlink(temp_path)



if __name__ == "__main__":
    unittest.main()
